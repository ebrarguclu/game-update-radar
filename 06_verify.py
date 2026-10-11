"""
Step 6: Verify findings by reading the reviews behind them
==========================================================

A number like "ad mentions +8 points" is only as good as the keyword list behind it.
This script turns each finding into something a human can check by reading.

Two modes:

  python 06_verify.py export
      Writes data/samples/verify.xlsx (raw text -> stays in data/, never published).
      One sheet per finding: matched reviews from the flagged release AND the release
      before it, so you can see whether the complaints really changed.
      One sheet per topic: a random sample of keyword matches to measure precision.
      Fill in the "correct" column (1 = keyword match is really about this topic, 0 = not)
      and write a few words in "about" (what is the complaint actually about?).

  python 06_verify.py score
      Reads your labels back and computes precision per sheet (with a ~95% interval).
      Writes results/verification.csv (aggregated numbers only, safe to publish).
"""

import importlib
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

analyze = importlib.import_module("03_analyze")   # reuse load(), TOPICS, release_table() - one source of truth

OUT_FILE = Path("data/samples/verify.xlsx")
SCORE_FILE = Path("results/verification.csv")
PER_GROUP = 25        # reviews per release in a finding sheet
PER_TOPIC = 40        # reviews per topic in a precision sheet
SEED = 42             # fixed seed: the same sample every run (reproducible)

# Findings from the README. "topic" = which keyword list produced the finding (None = all low-star reviews).
FINDINGS = [
    {"sheet": "RK 28796 ads",     "game": "royal_kingdom",  "version": "28796", "topic": "ads"},
    {"sheet": "RM 37281 ads",     "game": "royal_match",    "version": "37281", "topic": "ads"},
    {"sheet": "PF 0.24.3 bugs",   "game": "pixel_flow",     "version": "0.24.3", "topic": "bugs"},
    {"sheet": "MV 6969 low stars", "game": "match_villains", "version": "6969",  "topic": None},
]

HEADER = ["row", "group", "game", "version", "stars", "date", "matched words", "review", "correct", "about"]
ORANGE = PatternFill("solid", fgColor="FDE3D6")


def sample(d: pd.DataFrame, n: int, rng) -> pd.DataFrame:
    return d.iloc[rng.choice(len(d), size=min(n, len(d)), replace=False)] if len(d) else d


def matched_words(text: str, pattern: str | None) -> str:
    """Show WHICH words triggered the match - the fastest way to spot false positives."""
    if pattern is None:
        return ""
    return ", ".join(sorted({m.group(0) for m in re.finditer(pattern, text)}))


def previous_version(game: str, df: pd.DataFrame, version: str) -> str | None:
    t = analyze.release_table(game, df)
    idx = t.index[t["version"] == version]
    return t.loc[idx[0] - 1, "version"] if len(idx) and idx[0] > 0 else None


def write_sheet(wb, name: str, rows: pd.DataFrame, pattern: str | None):
    ws = wb.create_sheet(name)
    ws.append(HEADER)
    for i, r in enumerate(rows.itertuples(), start=1):
        version = None if pd.isna(r.version) else r.version   # reviews without a version -> empty cell (Excel can't store pandas' <NA>)
        ws.append([i, r.group, r.game, version, r.score, r.at.date(),
                   matched_words(r.text, pattern), r.content, None, None])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = ORANGE
    for col, width in zip("ABCDEFGHIJ", [5, 12, 15, 9, 6, 11, 18, 90, 9, 30]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        row[7].alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"1,0"', allow_blank=True)   # dropdown: only 1 or 0
    ws.add_data_validation(dv)
    dv.add(f"I2:I{ws.max_row}")
    ws.freeze_panes = "A2"


def export():
    rng = np.random.default_rng(SEED)
    wb = Workbook()
    wb.remove(wb.active)
    cache = {}

    # 1) Finding sheets: flagged release vs the release before it.
    for f in FINDINGS:
        df = cache.setdefault(f["game"], analyze.load(f["game"]))
        prev = previous_version(f["game"], df, f["version"])
        if prev is None:
            print(f"! {f['sheet']}: version {f['version']} not found or has no previous release, skipped")
            continue
        parts = []
        for label, v in [("before " + prev, prev), ("after " + f["version"], f["version"])]:
            d = df[df["version"] == v]
            d = d[d[f"t_{f['topic']}"]] if f["topic"] else d[d["score"] <= 2]
            parts.append(sample(d, PER_GROUP, rng).assign(group=label, game=f["game"]))
            print(f"{f['sheet']:20s} {label:16s} {len(d):5d} matching reviews, sampled {min(PER_GROUP, len(d))}")
        pattern = analyze.TOPICS[f["topic"]] if f["topic"] else None
        write_sheet(wb, f["sheet"], pd.concat(parts), pattern)

    # 2) Precision sheets: random keyword matches per topic, pooled over all games.
    all_games = pd.concat([cache.get(g, analyze.load(g)).assign(game=g) for g in analyze.GAMES])
    topics = dict(analyze.TOPICS, quit=analyze.QUIT_WORDS)
    all_games["t_quit"] = all_games["mentions_quit"]
    for topic, pattern in topics.items():
        d = all_games[all_games[f"t_{topic}"]]
        if d.empty:
            print(f"! precision {topic}: no matches, skipped")
            continue
        write_sheet(wb, f"precision {topic}", sample(d, PER_TOPIC, rng).assign(group="random match"), pattern)
        print(f"precision {topic:12s} {len(d):6d} matches, sampled {min(PER_TOPIC, len(d))}")

    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_FILE)
    print(f"\nSaved {OUT_FILE}. Fill in 'correct' (1/0) and 'about', save, then run: python 06_verify.py score")


def score():
    wb = load_workbook(OUT_FILE)
    rows = []
    for ws in wb.worksheets:
        df = pd.DataFrame(ws.iter_rows(min_row=2, max_col=len(HEADER), values_only=True), columns=HEADER)
        for group, g in df.groupby("group"):
            lab = pd.to_numeric(g["correct"], errors="coerce").dropna()
            n = len(lab)
            p = lab.mean() if n else np.nan
            se = np.sqrt(p * (1 - p) / n) if n else np.nan
            rows.append({"sheet": ws.title, "group": group, "labeled": n, "of": len(g),
                         "precision_pct": round(100 * p, 1) if n else np.nan,
                         "low_pct": round(100 * max(p - 2 * se, 0), 1) if n else np.nan,
                         "high_pct": round(100 * min(p + 2 * se, 1), 1) if n else np.nan})
    # Category codes written in "about" (ingame / misleading / noads / other): how are they split per group?
    for ws in wb.worksheets:
        df = pd.DataFrame(ws.iter_rows(min_row=2, max_col=len(HEADER), values_only=True), columns=HEADER)
        codes = df["about"].astype(str).str.strip().str.lower()
        if codes.isin(["ingame", "misleading", "noads"]).any():
            print(f"\n{ws.title}: what the keyword matches are really about (% of labeled rows)")
            print((pd.crosstab(df["group"], codes, normalize="index") * 100).round(0).to_string())

    out = pd.DataFrame(rows)
    SCORE_FILE.parent.mkdir(exist_ok=True)
    out.to_csv(SCORE_FILE, index=False)
    print(out.to_string(index=False))
    print(f"\nSaved {SCORE_FILE} (numbers only, no review text).")


if __name__ == "__main__":
    mode = sys.argv[1] if len(sys.argv) > 1 else ""
    if mode == "export":
        export()
    elif mode == "score":
        score()
    else:
        print(__doc__)
